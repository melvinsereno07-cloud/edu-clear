import io
import os
import sqlite3
import uuid
from datetime import datetime
from functools import wraps
from pathlib import Path
from urllib.parse import quote

from flask import (
    Flask,
    flash,
    g,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

try:
    from openpyxl import Workbook, load_workbook
except Exception:  # pragma: no cover
    Workbook = None
    load_workbook = None

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "see_oms.db"
UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "dev-secret-change-me")
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024


SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    full_name TEXT NOT NULL,
    position TEXT,
    username TEXT UNIQUE NOT NULL,
    password_hash TEXT NOT NULL,
    role TEXT NOT NULL,
    grade TEXT,
    section TEXT,
    status TEXT DEFAULT 'Active',
    qr_token TEXT UNIQUE,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS announcements (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    title TEXT NOT NULL,
    content TEXT NOT NULL,
    image_path TEXT,
    doc_path TEXT,
    pinned INTEGER DEFAULT 0,
    created_by INTEGER,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS fees (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    title TEXT NOT NULL,
    amount REAL NOT NULL,
    created_by INTEGER,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS payments (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id INTEGER NOT NULL,
    fee_id INTEGER NOT NULL,
    amount REAL NOT NULL,
    status TEXT NOT NULL,
    paid_at TEXT NOT NULL,
    receipt_no TEXT NOT NULL,
    recorded_by INTEGER,
    FOREIGN KEY(user_id) REFERENCES users(id),
    FOREIGN KEY(fee_id) REFERENCES fees(id)
);
CREATE TABLE IF NOT EXISTS attendance_events (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    title TEXT NOT NULL,
    event_date TEXT NOT NULL,
    created_by INTEGER,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS attendance_logs (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id INTEGER NOT NULL,
    event_id INTEGER NOT NULL,
    method TEXT NOT NULL,
    status TEXT NOT NULL,
    timestamp TEXT NOT NULL,
    recorded_by INTEGER,
    FOREIGN KEY(user_id) REFERENCES users(id),
    FOREIGN KEY(event_id) REFERENCES attendance_events(id)
);
CREATE TABLE IF NOT EXISTS org_structure (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    full_name TEXT NOT NULL,
    position TEXT NOT NULL,
    image_path TEXT
);
"""


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = sqlite3.connect(DB_PATH)
    db.executescript(SCHEMA_SQL)
    admin = db.execute("SELECT id FROM users WHERE username=?", ("melvin.sereno.admin",)).fetchone()
    if not admin:
        db.execute(
            """INSERT INTO users (full_name, position, username, password_hash, role, grade, section, status, qr_token)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                "Melvin Sereno",
                "System Administrator",
                "melvin.sereno.admin",
                generate_password_hash("baby@nerio2007"),
                "Admin",
                "N/A",
                "N/A",
                "Active",
                str(uuid.uuid4()),
            ),
        )
    db.commit()
    db.close()


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login"))
        return view(*args, **kwargs)

    return wrapped


def role_required(*roles):
    def dec(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if session.get("role") not in roles:
                flash("Unauthorized access.", "danger")
                return redirect(url_for("dashboard"))
            return view(*args, **kwargs)

        return wrapped

    return dec


def save_upload(file_obj):
    if not file_obj or not file_obj.filename:
        return None
    filename = f"{uuid.uuid4()}_{secure_filename(file_obj.filename)}"
    path = UPLOAD_DIR / filename
    file_obj.save(path)
    return filename


def get_qr_url(token):
    return f"https://quickchart.io/qr?size=180&text={quote(token)}"


init_db()


@app.context_processor
def inject_helpers():
    return {"qr_url": get_qr_url, "now": datetime.now}


@app.route("/", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form["username"].strip()
        password = request.form["password"]
        db = get_db()
        user = db.execute("SELECT * FROM users WHERE username=? AND status='Active'", (username,)).fetchone()
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            session["role"] = user["role"]
            session["full_name"] = user["full_name"]
            return redirect(url_for("dashboard"))
        flash("Invalid username or password.", "danger")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/dashboard")
@login_required
def dashboard():
    db = get_db()
    role = session["role"]
    uid = session["user_id"]

    total_students = db.execute("SELECT COUNT(*) c FROM users WHERE role='Student'").fetchone()["c"]
    total_ann = db.execute("SELECT COUNT(*) c FROM announcements").fetchone()["c"]

    if role == "Student":
        financial = db.execute(
            """SELECT COALESCE(SUM(p.amount),0) paid, COALESCE(SUM(f.amount),0) due
               FROM fees f LEFT JOIN payments p ON p.fee_id=f.id AND p.user_id=?""",
            (uid,),
        ).fetchone()
        attendance = db.execute("SELECT COUNT(*) c FROM attendance_logs WHERE user_id=?", (uid,)).fetchone()["c"]
    else:
        financial = db.execute(
            "SELECT COALESCE(SUM(amount),0) paid FROM payments"
        ).fetchone()
        attendance = db.execute("SELECT COUNT(*) c FROM attendance_logs").fetchone()["c"]

    announcements = db.execute(
        "SELECT * FROM announcements ORDER BY pinned DESC, created_at DESC LIMIT 5"
    ).fetchall()

    return render_template(
        "dashboard.html",
        role=role,
        financial=financial,
        attendance=attendance,
        total_students=total_students,
        total_ann=total_ann,
        announcements=announcements,
    )


@app.route("/users", methods=["GET", "POST"])
@login_required
@role_required("Admin")
def users():
    db = get_db()
    if request.method == "POST":
        db.execute(
            """INSERT INTO users (full_name, position, username, password_hash, role, grade, section, status, qr_token)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                request.form["full_name"],
                request.form.get("position", ""),
                request.form["username"],
                generate_password_hash(request.form["password"]),
                request.form["role"],
                request.form.get("grade", ""),
                request.form.get("section", ""),
                request.form.get("status", "Active"),
                str(uuid.uuid4()),
            ),
        )
        db.commit()
        flash("Account created.", "success")
        return redirect(url_for("users"))
    users_data = db.execute("SELECT * FROM users ORDER BY created_at DESC").fetchall()
    return render_template("users.html", users=users_data)


@app.post("/users/<int:user_id>/delete")
@login_required
@role_required("Admin")
def delete_user(user_id):
    db = get_db()
    db.execute("DELETE FROM users WHERE id=?", (user_id,))
    db.commit()
    flash("User deleted.", "success")
    return redirect(url_for("users"))


@app.route("/users/bulk-upload", methods=["POST"])
@login_required
@role_required("Admin")
def bulk_upload_users():
    if load_workbook is None:
        flash("openpyxl is required for Excel upload.", "danger")
        return redirect(url_for("users"))

    excel = request.files.get("excel_file")
    if not excel:
        flash("Please upload a file.", "danger")
        return redirect(url_for("users"))

    wb = load_workbook(excel)
    ws = wb.active
    db = get_db()
    inserted = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not row[2] or not row[3]:
            continue
        full_name, position, username, password, role, grade, section, status = row[:8]
        try:
            db.execute(
                """INSERT INTO users (full_name, position, username, password_hash, role, grade, section, status, qr_token)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    str(full_name),
                    str(position or ""),
                    str(username),
                    generate_password_hash(str(password)),
                    str(role or "Student"),
                    str(grade or ""),
                    str(section or ""),
                    str(status or "Active"),
                    str(uuid.uuid4()),
                ),
            )
            inserted += 1
        except sqlite3.IntegrityError:
            continue
    db.commit()
    flash(f"Bulk upload completed. Added {inserted} users.", "success")
    return redirect(url_for("users"))


@app.route("/users/template")
@login_required
@role_required("Admin")
def user_template():
    if Workbook is None:
        flash("openpyxl is required for Excel template.", "danger")
        return redirect(url_for("users"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    ws.append(["Full Name", "Position", "Username", "Password", "Role", "Grade", "Section", "Status"])
    ws.append(["Juan Dela Cruz", "Treasurer", "juan.student", "secure123", "Student", "11", "Rizal", "Active"])
    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    return send_file(mem, as_attachment=True, download_name="bulk_user_template.xlsx")


@app.route("/announcements", methods=["GET", "POST"])
@login_required
def announcements():
    db = get_db()
    if request.method == "POST":
        if session["role"] != "Admin":
            flash("Only Admin can create announcements.", "danger")
            return redirect(url_for("announcements"))
        image_file = save_upload(request.files.get("image"))
        doc_file = save_upload(request.files.get("document"))
        db.execute(
            "INSERT INTO announcements (title, content, image_path, doc_path, pinned, created_by) VALUES (?, ?, ?, ?, ?, ?)",
            (
                request.form["title"],
                request.form["content"],
                image_file,
                doc_file,
                1 if request.form.get("pinned") else 0,
                session["user_id"],
            ),
        )
        db.commit()
        flash("Announcement posted.", "success")
        return redirect(url_for("announcements"))

    items = db.execute("SELECT * FROM announcements ORDER BY pinned DESC, created_at DESC").fetchall()
    return render_template("announcements.html", announcements=items)


@app.post("/announcements/<int:ann_id>/pin")
@login_required
@role_required("Admin")
def pin_announcement(ann_id):
    db = get_db()
    db.execute("UPDATE announcements SET pinned = CASE WHEN pinned=1 THEN 0 ELSE 1 END WHERE id=?", (ann_id,))
    db.commit()
    return redirect(url_for("announcements"))


@app.route("/financial", methods=["GET", "POST"])
@login_required
def financial():
    db = get_db()
    role = session["role"]

    if request.method == "POST" and role == "Admin":
        db.execute(
            "INSERT INTO fees (title, amount, created_by) VALUES (?, ?, ?)",
            (request.form["title"], float(request.form["amount"]), session["user_id"]),
        )
        db.commit()
        flash("Fee created.", "success")
        return redirect(url_for("financial"))

    fees = db.execute("SELECT * FROM fees ORDER BY created_at DESC").fetchall()
    students = db.execute("SELECT id, full_name FROM users WHERE role='Student' AND status='Active'").fetchall()
    records = db.execute(
        """SELECT p.*, u.full_name, f.title FROM payments p
           JOIN users u ON p.user_id=u.id JOIN fees f ON p.fee_id=f.id
           ORDER BY p.paid_at DESC"""
    ).fetchall()
    my_records = []
    if role == "Student":
        my_records = db.execute(
            "SELECT p.*, f.title FROM payments p JOIN fees f ON p.fee_id=f.id WHERE p.user_id=? ORDER BY p.paid_at DESC",
            (session["user_id"],),
        ).fetchall()

    return render_template("financial.html", role=role, fees=fees, students=students, records=records, my_records=my_records)


@app.post("/financial/pay")
@login_required
@role_required("Admin", "Manager")
def pay_fee():
    db = get_db()
    amount = float(request.form["amount"])
    fee = db.execute("SELECT * FROM fees WHERE id=?", (request.form["fee_id"],)).fetchone()
    status = "Fully Paid" if amount >= fee["amount"] else "Partially Paid"
    paid_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    receipt_no = f"REC-{uuid.uuid4().hex[:8].upper()}"
    db.execute(
        "INSERT INTO payments (user_id, fee_id, amount, status, paid_at, receipt_no, recorded_by) VALUES (?, ?, ?, ?, ?, ?, ?)",
        (request.form["student_id"], request.form["fee_id"], amount, status, paid_at, receipt_no, session["user_id"]),
    )
    db.commit()
    flash(f"Payment recorded with receipt {receipt_no}", "success")
    return redirect(url_for("receipt", receipt_no=receipt_no))


@app.route("/receipt/<receipt_no>")
@login_required
def receipt(receipt_no):
    db = get_db()
    row = db.execute(
        """SELECT p.*, u.full_name, f.title FROM payments p
           JOIN users u ON p.user_id=u.id JOIN fees f ON p.fee_id=f.id WHERE receipt_no=?""",
        (receipt_no,),
    ).fetchone()
    return render_template("receipt.html", row=row)


@app.route("/attendance", methods=["GET", "POST"])
@login_required
def attendance():
    db = get_db()
    role = session["role"]
    if request.method == "POST" and role in ["Admin", "Manager"]:
        db.execute(
            "INSERT INTO attendance_events (title, event_date, created_by) VALUES (?, ?, ?)",
            (request.form["title"], request.form["event_date"], session["user_id"]),
        )
        db.commit()
        flash("Attendance event created.", "success")
        return redirect(url_for("attendance"))

    events = db.execute("SELECT * FROM attendance_events ORDER BY event_date DESC").fetchall()
    logs = db.execute(
        """SELECT a.*, u.full_name, e.title event_title FROM attendance_logs a
        JOIN users u ON a.user_id=u.id JOIN attendance_events e ON a.event_id=e.id
        ORDER BY timestamp DESC"""
    ).fetchall()
    my_logs = []
    my_qr = None
    if role == "Student":
        my_logs = db.execute(
            """SELECT a.*, e.title event_title FROM attendance_logs a
               JOIN attendance_events e ON a.event_id=e.id WHERE a.user_id=? ORDER BY timestamp DESC""",
            (session["user_id"],),
        ).fetchall()
        my_qr = db.execute("SELECT qr_token FROM users WHERE id=?", (session["user_id"],)).fetchone()["qr_token"]
    return render_template("attendance.html", role=role, events=events, logs=logs, my_logs=my_logs, my_qr=my_qr)


@app.post("/attendance/record")
@login_required
@role_required("Admin", "Manager")
def record_attendance():
    db = get_db()
    token = request.form.get("qr_token", "").strip()
    method = "QR Scan" if token else "Manual"
    if token:
        student = db.execute("SELECT id FROM users WHERE qr_token=?", (token,)).fetchone()
    else:
        student = db.execute("SELECT id FROM users WHERE id=?", (request.form.get("student_id"),)).fetchone()
    if not student:
        flash("Invalid student/QR token.", "danger")
        return redirect(url_for("attendance"))

    db.execute(
        "INSERT INTO attendance_logs (user_id, event_id, method, status, timestamp, recorded_by) VALUES (?, ?, ?, ?, ?, ?)",
        (
            student["id"],
            request.form["event_id"],
            method,
            request.form.get("status", "Present"),
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            session["user_id"],
        ),
    )
    db.commit()
    flash("Attendance recorded.", "success")
    return redirect(url_for("attendance"))


@app.route("/org-structure", methods=["GET", "POST"])
@login_required
def org_structure():
    db = get_db()
    if request.method == "POST":
        if session["role"] != "Admin":
            flash("Only Admin can update organization structure.", "danger")
            return redirect(url_for("org_structure"))
        image_file = save_upload(request.files.get("image"))
        db.execute(
            "INSERT INTO org_structure (full_name, position, image_path) VALUES (?, ?, ?)",
            (request.form["full_name"], request.form["position"], image_file),
        )
        db.commit()
        flash("Officer card added.", "success")
        return redirect(url_for("org_structure"))

    cards = db.execute("SELECT * FROM org_structure ORDER BY position").fetchall()
    return render_template("org_structure.html", cards=cards)


@app.route("/uploads/<path:filename>")
@login_required
def uploads(filename):
    return send_file(UPLOAD_DIR / filename)


@app.route("/reports/financial")
@login_required
@role_required("Admin", "Manager")
def financial_report():
    if Workbook is None:
        flash("openpyxl is required for Excel reports.", "danger")
        return redirect(url_for("financial"))
    db = get_db()
    rows = db.execute(
        """SELECT u.full_name, f.title, p.amount, p.status, p.paid_at, p.receipt_no FROM payments p
           JOIN users u ON p.user_id=u.id JOIN fees f ON p.fee_id=f.id ORDER BY p.paid_at DESC"""
    ).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Financial Report"
    ws.append(["Student", "Fee", "Amount", "Status", "Paid At", "Receipt"])
    for r in rows:
        ws.append([r["full_name"], r["title"], r["amount"], r["status"], r["paid_at"], r["receipt_no"]])
    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    return send_file(mem, as_attachment=True, download_name="financial_report.xlsx")


@app.route("/reports/attendance")
@login_required
@role_required("Admin", "Manager")
def attendance_report():
    if Workbook is None:
        flash("openpyxl is required for Excel reports.", "danger")
        return redirect(url_for("attendance"))
    db = get_db()
    rows = db.execute(
        """SELECT u.full_name, e.title, a.method, a.status, a.timestamp FROM attendance_logs a
           JOIN users u ON a.user_id=u.id JOIN attendance_events e ON a.event_id=e.id
           ORDER BY a.timestamp DESC"""
    ).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Report"
    ws.append(["Student", "Event", "Method", "Status", "Timestamp"])
    for r in rows:
        ws.append([r["full_name"], r["title"], r["method"], r["status"], r["timestamp"]])
    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    return send_file(mem, as_attachment=True, download_name="attendance_report.xlsx")


if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=5000, debug=True)
